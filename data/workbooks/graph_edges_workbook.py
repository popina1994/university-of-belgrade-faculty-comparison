from openpyxl import Workbook
from data.tables.graph_edge import GraphEdge
from data.workbooks.works_workbook import WorkTypes
from utilities.global_setup import DATA_PATH
GRAPH_EDGES_WOS_FILE_NAME = DATA_PATH + r"\work\graph_edges_wos.xlsx"
GRAPH_EDGES_SCOPUS_FILE_NAME = DATA_PATH + r"\work\graph_edges_scopus.xlsx"
GRAPH_EDGES_TMP_FILE_NAME = DATA_PATH + r"\work\graph_edges_tmp.xlsx"
GRAPH_EDGES_FILE_NAMES = [GRAPH_EDGES_WOS_FILE_NAME, GRAPH_EDGES_SCOPUS_FILE_NAME, GRAPH_EDGES_TMP_FILE_NAME]
GRAPH_EDGES_FRACT_WOS_FILE_NAME = DATA_PATH + r"\work\graph_edges_fract_wos.xlsx"
GRAPH_EDGES_FRACT_SCOPUS_FILE_NAME = DATA_PATH + r"\work\graph_edges_fract_scopus.xlsx"
GRAPH_EDGES_FRACT_TMP_FILE_NAME = DATA_PATH + r"\work\graph_edges_fract_tmp.xlsx"
GRAPH_EDGES_FRACT_FILE_NAMES = [GRAPH_EDGES_FRACT_WOS_FILE_NAME, GRAPH_EDGES_FRACT_SCOPUS_FILE_NAME,
                                GRAPH_EDGES_FRACT_TMP_FILE_NAME]
GRAPH_EDGES_SHEET = "Veze"


class GraphEdgesWorkbook:
    def __init__(self, work_book_type: WorkTypes, is_fraction: bool):
        self.work_book = Workbook()
        self.work_book.remove(self.work_book.active)
        self.sheet = self.work_book.create_sheet(GRAPH_EDGES_SHEET)
        source_file_names = GRAPH_EDGES_FRACT_FILE_NAMES if is_fraction else GRAPH_EDGES_FILE_NAMES
        self.file_name = source_file_names[work_book_type]
        GraphEdge.write_headers_to_sheet(self.sheet)
        self.row = 2

    def save_graph_edge(self, graph_edge: GraphEdge):
        graph_edge.write_to_sheet(self.sheet, self.row)
        self.row += 1

    def save(self):
        self.work_book.save(self.file_name)
