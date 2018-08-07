from openpyxl import Workbook

from data.tables.author import Author
from data.workbooks.works_workbook import WorkTypes
from utilities.global_setup import DATA_PATH

ALL_AUTHORS_WOS_FILE_NAME = DATA_PATH + r"\people\authors_all_wos.xlsx"
ALL_AUTHORS_SCOPUS_FILE_NAME = DATA_PATH + r"\people\authors_all_scopus.xlsx"
ALL_AUTHORS_SHEET = "Svi"

ALL_AUTHORS_FILE_NAMES = [ALL_AUTHORS_WOS_FILE_NAME, ALL_AUTHORS_SCOPUS_FILE_NAME]


class AuthorsAllWorkBook:
    def __init__(self, work_book_type: WorkTypes):
        self.work_book = Workbook()
        self.work_book.remove(self.work_book.active)
        self.sheet = self.work_book.create_sheet(ALL_AUTHORS_SHEET)
        self.file_name = ALL_AUTHORS_FILE_NAMES[work_book_type]
        Author.write_headers_to_sheet(self.sheet)
        self.row = 2

    def save_author(self, author: Author):
        author.write_to_sheet(self.sheet, self.row)
        self.row += 1

    def save(self):
        self.work_book.save(self.file_name)
